"""把单个用户在原始数据 / 查询结果中的所有记录导出到一个 Excel。

每个 CSV 一个 sheet，首张 sheet 为目录，列出表来源、用户行数、全表行数和字段数。
用法示例：
    python export_user_case.py --user-id 1000566
"""

import argparse
import codecs
import re
import sys
import unicodedata
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")


PROJECT_DIR = Path(__file__).resolve().parent
DEFAULT_RAW_DIR = Path.home() / "Downloads" / "vintage_data" / "原始数据"
DEFAULT_RESULT_DIR = Path.home() / "Downloads" / "vintage_data" / "查询结果"
DEFAULT_USER_ID = "1000566"

SOURCE_LABELS = {"raw": "原始数据", "result": "查询结果"}
USER_KEYS = ("user_id", "borrower_id")

SHEETS = {
    "01": ("01_借款申请表", "借据主表：申请状态、放款时间、挂牌 ID"),
    "02": ("02_放款信息表", "挂牌主表：产品、期限、金额、利率、成交信息"),
    "03": ("03_债务快照表", "债务快照原始表：债务逐期状态与金额"),
    "04": ("04_风控事件表", "风控事件任务：一行一次风控流程"),
    "05": ("05_额度申请表", "额度申请：一行一次额度申请"),
    "06": ("06_额度日志表", "额度变动流水：占用 / 释放 / 过期"),
    "07": ("07_观察日参数表", "全局观察日参数"),
    "08": ("08_债务快照子集", "由 03 按挂牌 × 快照日聚合的债务快照"),
    "09": ("09_借款申请明细", "申请级明细：申请序列、是否首申、策略时间"),
    "10": ("10_借款债务明细", "申请 + 债务合并，一行一笔借据"),
    "11": ("11_放款成功表", "仅放款成功的借据，含放款月与 New / Old"),
    "12": ("12_戳额表", "放款对应的额度戳额记录"),
    "13": ("13_借款级Vintage宽表", "借据级 Vintage 宽表：标签 + MOB 5+ / 30+ 标记"),
    "14": ("14_放款样本表", "最终放款样本，一行一笔借据"),
}

ILLEGAL_XML_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

FONT_NAME = "微软雅黑"
DATA_FONT = Font(name=FONT_NAME, size=10)
BOLD_FONT = Font(name=FONT_NAME, size=10, bold=True)
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


def load_table(path):
    """读取 CSV，兼容带 BOM 的 UTF-8 与 GB18030。"""
    with path.open("rb") as handle:
        has_bom = handle.read(3) == codecs.BOM_UTF8
    if has_bom:
        encoding = "utf-8-sig"
    else:
        try:
            path.open("rb").read().decode("utf-8")
            encoding = "utf-8"
        except UnicodeDecodeError:
            encoding = "gb18030"
    df = pd.read_csv(path, encoding=encoding, encoding_errors="replace", low_memory=False)
    df.columns = [str(col).lstrip("﻿").strip() for col in df.columns]
    return df, encoding


def filter_user(df, user_id):
    key = next((col for col in USER_KEYS if col in df.columns), None)
    if key is None:
        return df.copy(), None
    values = (
        df[key]
        .astype("string")
        .str.strip()
        .str.replace(r"\.0+$", "", regex=True)
    )
    mask = (values == user_id).fillna(False).astype(bool)
    return df.loc[mask].copy(), key


def sanitize(df):
    for column in df.columns:
        if df[column].dtype == object:
            df[column] = df[column].map(
                lambda value: ILLEGAL_XML_CHARS.sub("", value)
                if isinstance(value, str)
                else value
            )
    return df


def display_width(value):
    if pd.isna(value):
        return 0
    return sum(
        2 if unicodedata.east_asian_width(char) in ("W", "F") else 1
        for char in str(value)
    )


def style_sheet(worksheet, frame, kind):
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_properties.tabColor = "1F4E78" if kind == "raw" else "548235"
    for row in worksheet.iter_rows():
        for cell in row:
            cell.font = DATA_FONT
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    worksheet.row_dimensions[1].height = 24
    for index, column in enumerate(frame.columns, start=1):
        width = display_width(column)
        for value in frame[column].head(200):
            width = max(width, display_width(value))
        worksheet.column_dimensions[get_column_letter(index)].width = min(
            max(width + 2, 10), 50
        )


def style_index(worksheet, row_count):
    for row in worksheet.iter_rows():
        for cell in row:
            cell.font = DATA_FONT
    worksheet["A1"].font = BOLD_FONT
    worksheet["A2"].font = BOLD_FONT
    worksheet["A3"].font = BOLD_FONT
    worksheet["B1"].number_format = "@"
    header_row = 5
    for cell in worksheet[header_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    worksheet.row_dimensions[header_row].height = 24
    worksheet.freeze_panes = "A6"
    worksheet.auto_filter.ref = (
        f"A{header_row}:J{header_row + row_count}"
    )
    widths = {"A": 6, "B": 22, "C": 10, "D": 46, "E": 10, "F": 12,
              "G": 8, "H": 12, "I": 44, "J": 24}
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width


def collect_tables(raw_dir, result_dir):
    tables = {}
    for kind, folder in (("raw", raw_dir), ("result", result_dir)):
        if not folder.is_dir():
            raise FileNotFoundError(f"文件夹不存在：{folder}")
        for path in folder.glob("*.csv"):
            prefix = path.name[:2]
            if prefix in tables:
                raise ValueError(f"文件名前缀重复：{path.name}")
            tables[prefix] = (path, kind)
    missing = [f"{number:02d}" for number in range(1, 15) if f"{number:02d}" not in tables]
    if missing:
        raise FileNotFoundError(f"缺少以下文件：{', '.join(missing)}")
    return [(prefix, *tables[prefix]) for prefix in sorted(tables)]


def main(argv=None):
    parser = argparse.ArgumentParser(
        description="把单个用户的数据导出为一个 Excel，每张表一个 sheet。"
    )
    parser.add_argument("--user-id", default=DEFAULT_USER_ID, help="样本用户 ID")
    parser.add_argument("--raw-dir", type=Path, default=DEFAULT_RAW_DIR)
    parser.add_argument("--result-dir", type=Path, default=DEFAULT_RESULT_DIR)
    parser.add_argument("--output", type=Path, help="输出 xlsx 路径")
    args = parser.parse_args(argv)

    user_id = args.user_id.strip()
    output = args.output or PROJECT_DIR / f"vintage_user_case_{user_id}.xlsx"

    tables = collect_tables(args.raw_dir, args.result_dir)
    datasets = []
    index_rows = []

    for position, (prefix, path, kind) in enumerate(tables, start=1):
        sheet_name, description = SHEETS[prefix]
        frame, encoding = load_table(path)
        total_rows = len(frame)
        subset, key = filter_user(frame, user_id)
        subset = sanitize(subset)

        note = ""
        if key is None:
            note = "整表保留（无用户维度）"
        elif subset.empty:
            note = "该用户无记录"

        index_rows.append(
            {
                "序号": position,
                "Sheet": sheet_name,
                "来源": SOURCE_LABELS[kind],
                "源文件": path.name,
                "用户行数": len(subset),
                "全表行数": total_rows,
                "字段数": len(frame.columns),
                "用户键": key or "无",
                "说明": description,
                "备注": note,
            }
        )
        datasets.append((sheet_name, subset, kind))
        print(
            f"[{position:02d}/{len(tables)}] {path.name} | {encoding} | "
            f"用户 {len(subset):,} / 全表 {total_rows:,} 行"
        )

    index_frame = pd.DataFrame(index_rows)

    output.parent.mkdir(parents=True, exist_ok=True)
    try:
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            index_frame.to_excel(writer, sheet_name="00_目录", index=False, startrow=4)
            index_sheet = writer.sheets["00_目录"]
            index_sheet["A1"] = "样本用户ID"
            index_sheet["B1"] = user_id
            index_sheet["A2"] = "生成时间"
            index_sheet["B2"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            index_sheet["A3"] = "说明"
            index_sheet["B3"] = "仅保留该用户在各表中的记录；07 无用户维度，整表保留。"
            style_index(index_sheet, len(index_frame))

            for sheet_name, frame, kind in datasets:
                frame.to_excel(writer, sheet_name=sheet_name, index=False)
                style_sheet(writer.sheets[sheet_name], frame, kind)
    except PermissionError:
        print(f"无法写入 {output}：文件可能正在 Excel 中打开，请关闭后重试。")
        return 1

    user_rows = int(index_frame["用户行数"].sum())
    scoped = int((index_frame["用户键"] != "无").sum())
    print(f"\n已生成：{output}")
    print(
        f"样本用户 {user_id}：{len(datasets)} 张表，其中 {scoped} 张按用户过滤，"
        f"共 {user_rows:,} 行用户记录。"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
